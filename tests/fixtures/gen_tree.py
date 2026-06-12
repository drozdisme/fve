import openpyxl, os, time, csv
BASE = os.path.join(os.path.dirname(__file__), "tree")

def wb_with(path, labels, formula=None):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Calc"
    r = 1
    for lab, val in labels:
        ws.cell(row=r, column=1, value=lab); ws.cell(row=r, column=2, value=val); r += 1
    if formula:
        ws.cell(row=r, column=1, value="MS"); ws.cell(row=r, column=2, value=formula)
    wb.save(path)

# scattered structure across folders
wb_with(f"{BASE}/fuselage/section_655/stringers/stringer_loads_Rev_C.xlsx",
        [("BAY ID",12),("FEM LOAD CASE",3),("Nx",420.0),("Nxy",55.0),("MSstr_buck",1.2)], "=B3/(B4*B1)-1")
wb_with(f"{BASE}/fuselage/section_655/stringers/stringer_loads_Rev_D.xlsx",
        [("BAY ID",12),("FEM LOAD CASE",3),("Nx",430.0),("Nxy",58.0),("MSstr_buck",1.3)], "=B3/(B4*B1)-1")
wb_with(f"{BASE}/fuselage/section_655/frames/frame_station_120_v1.xlsx",
        [("FRAME",120),("STATION",655),("MOMENT OF INERTIA",4.2),("BENDING",900)], "=B3/B4-1")
wb_with(f"{BASE}/fuselage/section_655/frames/frame_station_120_v2.xlsx",
        [("FRAME",120),("STATION",655),("MOMENT OF INERTIA",4.4),("BENDING",880)], "=B3/B4-1")
wb_with(f"{BASE}/fuselage/section_676/fasteners/rivet_joint_R1.xlsx",
        [("RIVET","BACR15FV5"),("RIVET DIAMETER",4.8),("SHEAR",1200),("PITCH",20)], "=B3/1000-1")
wb_with(f"{BASE}/fuselage/section_676/fasteners/rivet_joint_R2.xlsx",
        [("RIVET","BACR15FV6"),("RIVET DIAMETER",5.0),("SHEAR",1300),("PITCH",22)], "=B3/1000-1")
wb_with(f"{BASE}/fuselage/section_676/fasteners/bolt_joint.xlsx",
        [("BOLT","HST12"),("BOLT DIAMETER",6.35),("TORQUE",9.0),("PRELOAD",5000),("SHANK",10)], "=B4/4000-1")

os.makedirs(f"{BASE}/wing/panels", exist_ok=True)
with open(f"{BASE}/wing/panels/panel_buckling.csv","w",newline="") as f:
    w = csv.writer(f); w.writerow(["PANEL","BUCKLING","sigma_cr","kc"]); w.writerow([1,"yes",250.0,4.0])

os.makedirs(f"{BASE}/misc", exist_ok=True)
open(f"{BASE}/misc/notes.txt","w").write("random engineer notes, no markers here\n")

# make Rev_D / v2 / R2 newer by mtime
now = time.time()
def older(p, days): os.utime(p, (now-days*86400, now-days*86400))
older(f"{BASE}/fuselage/section_655/stringers/stringer_loads_Rev_C.xlsx", 10)
older(f"{BASE}/fuselage/section_655/frames/frame_station_120_v1.xlsx", 10)
older(f"{BASE}/fuselage/section_676/fasteners/rivet_joint_R1.xlsx", 10)

n=sum(len(files) for _,_,files in os.walk(BASE))
print("tree files:", n)
