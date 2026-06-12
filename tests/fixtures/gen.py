import openpyxl, zipfile, struct, os, io

FX = os.path.dirname(__file__)

def holed():
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "S"
    ws["A1"]="load"; ws["B1"]=120
    ws["A2"]="cap";  ws["B2"]=200
    ws["A3"]="MS";   ws["B3"]="=IF(B1>0, B2/B1-1, 0)"
    ws["A4"]="look"; ws["B4"]="=VLOOKUP(B1, Tab!A1:B3, 2, FALSE)"
    t = wb.create_sheet("Tab")
    t["A1"]=100; t["B1"]=1.0
    t["A2"]=120; t["B2"]=1.2
    t["A3"]=150; t["B3"]=1.5
    wb.save(os.path.join(FX, "holed.xlsx"))

def img():
    src = os.path.join(FX, "beam.xlsx")
    dst = os.path.join(FX, "img.xlsx")
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108020000009077"
        "53de0000000c4944415408d76360000000020001f4716354000000004945"
        "4e44ae426082")
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zin.infolist():
            zo.writestr(it, zin.read(it.filename))
        zo.writestr("xl/media/image1.png", png)

def vstr(s):
    b = struct.pack("<I", len(s))
    for ch in s:
        b += struct.pack("<H", ord(ch))
    return b

def rec(rid, payload):
    out = b""
    if rid < 0x80:
        out += struct.pack("<B", rid)
    else:
        out += struct.pack("<B", (rid & 0x7F) | 0x80)
        out += struct.pack("<B", rid >> 7)
    n = len(payload)
    while True:
        b = n & 0x7F
        n >>= 7
        if n: out += struct.pack("<B", b | 0x80)
        else: out += struct.pack("<B", b); break
    return out + payload

def cell_real(col, val):
    return rec(0x05, struct.pack("<I", col) + struct.pack("<I", 0) + struct.pack("<d", val))

def ptg_ref(row, col):
    return struct.pack("<B", 0x24) + struct.pack("<I", row) + struct.pack("<H", col)

def fmla_add(col, val, r1, c1, r2, c2):
    rgce = ptg_ref(r1, c1) + ptg_ref(r2, c2) + struct.pack("<B", 0x03)
    payload = struct.pack("<I", col) + struct.pack("<I", 0) + struct.pack("<d", val)
    payload += struct.pack("<H", 0) + struct.pack("<I", len(rgce)) + rgce
    return rec(0x09, payload)

def xlsb():
    dst = os.path.join(FX, "mini.xlsb")
    wbk = rec(0x9C, struct.pack("<I", 0) + struct.pack("<I", 0) + vstr("") + vstr("Calc"))
    s = b""
    s += rec(0x00, struct.pack("<I", 0))            # row 0
    s += cell_real(0, 10.0)                          # A1 = 10
    s += cell_real(1, 32.0)                          # B1 = 32
    s += rec(0x00, struct.pack("<I", 1))            # row 1
    s += fmla_add(0, 42.0, 0, 0, 0, 1)               # A2 = A1+B1 cached 42
    ct = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="bin" ContentType="application/vnd.ms-excel.sheet.binary.macroEnabled.main"/></Types>')
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        zo.writestr("[Content_Types].xml", ct)
        zo.writestr("xl/workbook.bin", wbk)
        zo.writestr("xl/worksheets/sheet1.bin", s)

holed(); img(); xlsb()
print("fixtures:", sorted(os.listdir(FX)))
